"""Protocol, amount safety, API payload and interruption/recovery regression tests."""
import hashlib
import io
from contextlib import redirect_stderr
from decimal import Decimal
import json
import logging
from pathlib import Path
import random
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import httpx
from openai import OpenAI
from src.experiment.decisions import normalize_decision
from src.experiment.audit_decisions import audit_file
from src.experiment.llm import ExperimentPaused, MockChat, OpenAIChat, LLMResponse
from src.experiment.model import GameModel
from src.experiment.params import GameScenario, Params
from src.experiment.persistence import atomic_json
from src.experiment.run_main import CORE_FILES, build_scenarios, main, scrub

ROOT = Path(__file__).resolve().parents[1]


def setUpModule():
    (ROOT / "outputs").mkdir(exist_ok=True)
    logging.disable(logging.CRITICAL)


def tearDownModule():
    logging.disable(logging.NOTSET)


def params(rounds=1):
    result = Params(require_credentials=False)
    result.num_agents = 9
    result.proportions = [0.5]
    result.rounds = rounds
    result.max_tokens = 512
    result.retry_delay = 0
    result.validate()
    return result


class CountingMock(MockChat):
    def __init__(self, p, pause_at=None):
        super().__init__(p, "deepseek-flash")
        self.calls = []
        self.pause_at = pause_at

    def complete(self, messages, **kwargs):
        key = (kwargs["round_number"], kwargs["stage"], kwargs["agent"].unique_id)
        if len(self.calls) == self.pause_at:
            raise ExperimentPaused("injected_failure")
        self.calls.append(key)
        return super().complete(messages, **kwargs)


class ProtocolTests(unittest.TestCase):
    def test_output_examples_use_actual_neighbor_ids(self):
        p = params()
        for bdi in (False, True):
            p.report_bdi = bdi
            model = GameModel(params=p, gateway=CountingMock(p))
            for agent in model.agents:
                self.assertNotRegex(agent.system_prompt, r"N[1-4]_ID")
                for index, nid in enumerate(agent.neighbor_ids, 1):
                    self.assertIn(f'"{nid}":v{index}', agent.system_prompt)

    def test_incomplete_reply_pauses_before_action_and_cursor_commit(self):
        for finish in ("length", "content_filter", "tool_calls"):
            with self.subTest(finish=finish):
                p = params()
                events, snapshots = [], []
                class Incomplete(CountingMock):
                    def complete(self, messages, **kwargs):
                        result = super().complete(messages, **kwargs)
                        if len(self.calls) == 2:
                            result.finish_reason = finish
                            result.response_id = "incomplete-test"
                        return result
                gateway = Incomplete(p)
                model = GameModel(params=p, gateway=gateway,
                                  checkpoint=lambda m: snapshots.append(json.loads(json.dumps(m.to_state()))),
                                  event=lambda kind, **details: events.append((kind, details)))
                with self.assertRaises(ExperimentPaused) as exc:
                    model.run_model()
                self.assertEqual(exc.exception.reason, "output_truncated" if finish == "length" else "incomplete_response")
                self.assertEqual(model.next_agent, 1)
                self.assertEqual(len(model.token_usage), 1)
                self.assertEqual(model.get_agent(1).invested_amounts, [])
                self.assertEqual(exc.exception.details["response_id"], "incomplete-test")
                self.assertTrue(any(kind == "decision_not_accepted" for kind, _ in events))
                resumed_gateway = CountingMock(p)
                resumed = GameModel(params=p, gateway=resumed_gateway, state=snapshots[-1])
                expected = GameModel(params=p, gateway=CountingMock(p)).run_model()
                self.assertEqual(resumed.run_model(), expected)
                self.assertEqual(resumed_gateway.calls[0], (1, "investor", 1))

    def test_fixed_torus_seed_and_neighbors(self):
        for side in (3, 6, 7):
            p = params()
            scenario = GameScenario(num_agents=side**2, width=side, height=side, seed=51)
            a = GameModel(scenario, p, MockChat(p, "mock"))
            b = GameModel(scenario, p, MockChat(p, "mock"))
            self.assertEqual([x.type_restriction for x in a.agents], [x.type_restriction for x in b.agents])
            for agent in a.agents:
                self.assertEqual(len(set(agent.neighbor_ids)), 4)
                self.assertNotIn(agent.unique_id, agent.neighbor_ids)
                for nid in agent.neighbor_ids:
                    self.assertIn(agent.unique_id, a.get_agent(nid).neighbor_ids)
            self.assertEqual(a.get_agent(0).neighbor_ids, sorted([1, side-1, side, side*(side-1)]))

    def test_seed_derivation_is_stable_and_independent(self):
        p = params()
        p.iterations = 3
        first = build_scenarios(p)
        self.assertEqual(first, build_scenarios(p))
        self.assertEqual(len(set(s["seed"] for s in first)), 3)
        p.seed += 1
        self.assertNotEqual(first, build_scenarios(p))

    def test_paper_payoff_and_pair_return_direction(self):
        p = params()
        class Scripted:
            def complete(self, messages, *, agent, round_number, stage):
                values = ({i: 0.25 for i in agent.neighbor_ids} if stage == "investor" else
                          {i: agent.unique_id/100 for i in agent.neighbor_ids})
                return LLMResponse(json.dumps(values))
        model = GameModel(params=p, gateway=Scripted())
        rows, pairs, _ = model.run_model()
        self.assertEqual((len(rows), len(pairs)), (9, 36))
        self.assertAlmostEqual(sum(r["round_payoff"] for r in rows), 18)
        row = next(r for r in pairs if r["trustor_id"] == 0 and r["trustee_id"] == 1)
        self.assertEqual(row["returned_amount"], 0.01)  # B -> A, not A -> B (0).
        for agent in model.agents:
            self.assertEqual(agent.trustor_payoff, round(agent.total_received_return-agent.total_sent, 2))
            self.assertEqual(agent.balance, agent.round_payoff)

    def test_history_keeps_three_complete_rounds_and_current_stage(self):
        p = params(5)
        gateway = CountingMock(p)
        model = GameModel(params=p, gateway=gateway)
        model.run_model()
        agent = model.get_agent(0)
        messages = agent.messages(6, "CURRENT", 3)
        self.assertEqual(len(messages), 14)  # system + 3*(2 stages*2 messages) + current
        self.assertEqual([t["round"] for t in agent.history if t["round"] >= 3], [3,3,4,4,5,5])
        model2 = GameModel(params=p, gateway=CountingMock(p))
        p.history_rounds = None
        model2.run_model()
        self.assertEqual(len(model2.get_agent(0).messages(6, "CURRENT", None)), 22)

    def test_resume_mid_stages_skips_saved_calls_and_preserves_prompts(self):
        for failed_call in (2, 9, 13, 18, 27):
            with self.subTest(failed_call=failed_call), tempfile.TemporaryDirectory() as temp:
                p = params(2)
                reference = GameModel(params=p, gateway=CountingMock(p))
                expected = reference.run_model()
                path = Path(temp) / "checkpoint.json"
                first_gateway = CountingMock(p, pause_at=failed_call)
                first = GameModel(params=p, gateway=first_gateway,
                                  checkpoint=lambda m: atomic_json(path, m.to_state()))
                with self.assertRaises(ExperimentPaused):
                    first.run_model()
                state = json.loads(path.read_text(encoding="utf-8"))
                next_gateway = CountingMock(p)
                resumed = GameModel(GameScenario(**state["scenario"]), p, next_gateway, state=state)
                actual = resumed.run_model()
                self.assertEqual(actual, expected)
                self.assertEqual(len(first_gateway.calls)+len(next_gateway.calls), 36)
                self.assertEqual(len(set(first_gateway.calls+next_gateway.calls)), 36)

    def test_resume_after_stage_or_round_checkpoint_is_idempotent(self):
        for target in (("investor", 9, 1), ("trustee", 0, 1), ("trustee", 9, 1),
                       ("finalize", 0, 1), ("investor", 0, 2), ("complete", 0, 3)):
            p = params(2)
            reference = GameModel(params=p, gateway=CountingMock(p))
            expected = reference.run_model()
            captured = []
            def save(model):
                if not captured and (model.phase, model.next_agent, model.step) == target:
                    captured.append(json.loads(json.dumps(model.to_state())))
                    raise ExperimentPaused("injected_checkpoint_pause")
            first_gateway = CountingMock(p)
            first = GameModel(params=p, gateway=first_gateway, checkpoint=save)
            with self.assertRaises(ExperimentPaused):
                first.run_model()
            gateway = CountingMock(p)
            resumed = GameModel(GameScenario(**captured[0]["scenario"]), p, gateway, state=captured[0])
            self.assertEqual(resumed.run_model(), expected)
            self.assertEqual(len(first_gateway.calls)+len(gateway.calls), 36)


class DecisionTests(unittest.TestCase):
    neighbors = [1, 2, 3, 4]

    def norm(self, text, stage="investor", restriction="free", budget=5):
        return normalize_decision(text, self.neighbors, stage, restriction, budget)

    def test_real_archived_bdi_replies(self):
        cases = json.loads((ROOT / "test/fixtures/bdi_decisions_20261013.json").read_text(encoding="utf-8"))
        for case in cases:
            with self.subTest(source=case["source"], stage=case["stage"]):
                d = normalize_decision(case["content"], case["neighbor_ids"], case["stage"],
                                       case["restriction"], case["budget"])
                self.assertFalse(d.parse_failed)
                self.assertEqual(d.amounts, {int(k): v for k, v in case["expected"].items()})

    def test_aliases_nested_bdi_and_no_unrelated_dictionary(self):
        d = self.norm('Belief {"risk": 2} Decision {N1_ID:0.1,N2_ID:0.2,N3_ID:0.3,N4_ID:0.4}')
        self.assertEqual(d.amounts, {1:0.1,2:0.2,3:0.3,4:0.4})
        self.assertIn("neighbor_alias_resolved", d.corrections)
        d = self.norm('{"Belief":"test","Decision":{"1_ID":1,"N2_ID":1,"3":1,"4":1},"other":{"risk":5}}')
        self.assertEqual(d.amounts, {i:1 for i in self.neighbors})
        self.assertTrue(self.norm('{"risk":5}').parse_failed)
        self.assertTrue(self.norm('{"999":5}').parse_failed)
        self.assertTrue(self.norm('{bad1: 5}').parse_failed)

    def test_prose_does_not_infer_hypothetical_historical_or_unfinished_actions(self):
        for text in ('Belief: I will send 1 token to each neighbor.',
                     'Intention: If I return 4.0 tokens total, it may help.',
                     'Intention: I considered 1.0 each, but decided against it.',
                     'Intention: If I send neighbor 1:1, neighbor 2:1, neighbor 3:1, neighbor 4:1, I risk losses.',
                     'Intention: I will send 1 token to each neighbor.\nAction: {1:1,2:',
                     'Belief: Last round {1:1,2:1,3:1,4:1}.\nIntention: Final {1:0.5,2:',
                     'Intention: Send money to each neighbor.'):
            with self.subTest(text=text):
                self.assertTrue(self.norm(text, stage="trustee" if "return" in text else "investor").parse_failed)
        self.assertEqual(self.norm('Intention: I will send 0 tokens to each neighbor.').amounts, {i:0 for i in self.neighbors})
        self.assertFalse(self.norm('Intention: I will send 0 tokens to each neighbor.').parse_failed)

    def test_budget_scaling_preserves_proportions(self):
        d = self.norm('{1: 6, 2: 3, 3: 1, 4: 0}')
        self.assertEqual(d.amounts, {1:3.0, 2:1.5, 3:0.5, 4:0.0})
        self.assertIn("scaled_to_budget", d.corrections)

    def test_equal_payments_negative_missing_unknown_and_precision(self):
        d = self.norm('{1: -1, 2: 2.005, 99: 8}', restriction="constrained")
        self.assertEqual(d.amounts, {i:0.50 for i in self.neighbors})
        self.assertIn("negative_to_zero", d.corrections)
        self.assertIn("non_neighbor_ignored", d.corrections)
        self.assertIn("missing_neighbors_to_zero", d.corrections)
        self.assertEqual(self.norm('{1:1,2:2,3:3,4:4}',stage="trustee",budget=3).amounts,
                         {i:0.75 for i in self.neighbors})

    def test_unparseable_and_nonfinite(self):
        self.assertTrue(self.norm("sorry, no decision").parse_failed)
        self.assertEqual(self.norm("sorry, no decision").amounts, {i:0 for i in self.neighbors})
        d = self.norm('{"1": NaN, "2": Infinity, "3": true, "4": "abc"}')
        self.assertEqual(d.amounts, {i:0 for i in self.neighbors})

    def test_total_return_and_final_bdi_dictionary(self):
        self.assertEqual(self.norm('{"total_return": 1.03}', stage="trustee",budget=3).amounts,
                         {i:0.25 for i in self.neighbors})
        for text in ('{Return: 1.03}', '{"RETURN": 1.03}', 'Return: 1.03'):
            self.assertEqual(self.norm(text, stage="trustee", budget=3).amounts, {i:0.25 for i in self.neighbors})
        self.assertEqual(self.norm('explanation {1: 9} final {1:1,2:1,3:1,4:1}').amounts,
                         {i:1 for i in self.neighbors})

    def test_generated_decisions_always_obey_invariants(self):
        rng = random.Random(71)
        for _ in range(150):
            for stage in ("investor", "trustee"):
                restriction = rng.choice(["free", "constrained"])
                budget = Decimal(rng.randint(0,6000))/100 if stage == "trustee" else Decimal(5)
                raw = {str(i): rng.uniform(-10,100) for i in self.neighbors + [99]}
                d = self.norm(json.dumps(raw), stage, restriction, budget)
                values = [Decimal(str(v)) for v in d.amounts.values()]
                self.assertEqual(set(d.amounts), set(self.neighbors))
                self.assertTrue(all(v >= 0 and v == v.quantize(Decimal('.01')) for v in values))
                self.assertLessEqual(sum(values), budget)
                if stage == "trustee" or restriction == "constrained":
                    self.assertEqual(len(set(values)), 1)


class APITests(unittest.TestCase):
    def client(self, handler):
        return OpenAI(api_key="test-secret", base_url="https://test.invalid/v1",
                      http_client=httpx.Client(transport=httpx.MockTransport(handler)), max_retries=0)

    def response(self):
        return {"id":"test-response", "object":"chat.completion", "created":0, "model":"flash-test",
                "choices":[{"index":0,"message":{"role":"assistant","content":"{1:0,2:0,3:0,4:0}"},
                            "finish_reason":"stop"}],
                "usage":{"prompt_tokens":10,"completion_tokens":5,"total_tokens":15}}

    def request(self, gateway):
        return gateway.complete([{"role":"system","content":"test"},{"role":"user","content":"decide"}],
                                agent=SimpleNamespace(unique_id=0),round_number=1,stage="investor")

    def test_deepseek_thinking_disabled_and_generic_compatibility(self):
        bodies=[]
        def handler(request):
            bodies.append(json.loads(request.content))
            return httpx.Response(200,json=self.response())
        p=params()
        with self.client(handler) as client:
            gateway=OpenAIChat(p,"deepseek-flash",client)
            result=self.request(gateway)
            self.assertEqual(result.usage["total_tokens"],15)
            self.assertEqual(bodies[-1]["thinking"],{"type":"disabled"})
            p.thinking=True
            self.request(gateway)
            self.assertEqual(bodies[-1]["thinking"],{"type":"enabled"})
            self.assertNotIn("temperature",bodies[-1])
            self.request(OpenAIChat(p,"some-other-provider",client))
            self.assertNotIn("thinking",bodies[-1])

    def test_transient_retry_then_success_and_permanent_pause(self):
        calls=[]
        def handler(request):
            calls.append(request)
            return httpx.Response(429,json={"error":{"message":"busy","type":"rate_limit"}}) if len(calls)==1 else httpx.Response(200,json=self.response())
        with self.client(handler) as client:
            result=self.request(OpenAIChat(params(),"deepseek-flash",client))
            self.assertEqual((len(calls),result.attempts),(2,2))
        calls=[]
        def unauthorized(request):
            calls.append(request)
            return httpx.Response(401,json={"error":{"message":"no credit or auth","type":"auth"}})
        with self.client(unauthorized) as client:
            with self.assertRaises(ExperimentPaused):
                self.request(OpenAIChat(params(),"deepseek-flash",client))
        self.assertEqual(len(calls),1)

    def test_local_and_provider_context_limits_pause_without_zero_actions(self):
        calls=[]
        def handler(request):
            calls.append(request)
            return httpx.Response(400,json={"error":{"message":"maximum context length exceeded","type":"context_length_exceeded"}})
        p=params()
        with self.client(handler) as client:
            p.context_limit=520
            with self.assertRaises(ExperimentPaused) as exc:
                self.request(OpenAIChat(p,"deepseek-flash",client))
            self.assertEqual(exc.exception.reason,"context_budget_exceeded")
            self.assertEqual(calls,[])
            p.context_limit=131072
            with self.assertRaises(ExperimentPaused) as exc:
                self.request(OpenAIChat(p,"deepseek-flash",client))
            self.assertEqual(exc.exception.reason,"provider_context_exceeded")
            self.assertEqual(len(calls),1)


class CLITests(unittest.TestCase):
    def setUp(self):
        # Tests use a small explicit scenario even when the user's run defaults change.
        defaults = patch("src.experiment.run_main.Params", side_effect=lambda **_: params())
        defaults.start()
        self.addCleanup(defaults.stop)

    def test_truncated_reply_is_archived_and_resume_retries_only_pending_agent(self):
        with tempfile.TemporaryDirectory(dir=ROOT / "outputs") as temp:
            out = Path(temp) / "truncated"
            class Truncated(CountingMock):
                def complete(self, messages, **kwargs):
                    response = super().complete(messages, **kwargs)
                    if len(self.calls) == 2:
                        response.finish_reason = "length"
                        response.response_id = "truncated-reply"
                    return response
            with patch("src.experiment.run_main.MockChat", side_effect=lambda p, name, **_: Truncated(p)):
                self.assertEqual(main(["--mock", "--output-dir", str(out)]), 2)
            saved = json.loads((out / "checkpoint.json").read_text(encoding="utf-8"))
            self.assertEqual(saved["pause"]["reason"], "output_truncated")
            self.assertEqual(saved["pause"]["details"]["response_id"], "truncated-reply")
            self.assertEqual(saved["current"]["next_agent"], 1)
            self.assertTrue(saved["pause"]["details"]["content"])
            self.assertEqual(main(["--resume", str(out)]), 0)
            self.assertEqual(json.loads((out / "summary.json").read_text())["accepted_calls"], 18)

    def test_audit_reports_intention_without_rewriting_source(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "dialogue.json"
            turn = {"content": "{1_ID:0.5,2_ID:0.5,3_ID:0.5,4_ID:0.5}",
                    "prompt": "decide", "executed_decision": {str(i):0 for i in (1,2,3,4)},
                    "parse_failed": True, "finish_reason": "stop"}
            atomic_json(path, {"dialogs":{"round_1":{"agent_0_free":{"as_investor":turn}}}})
            before = path.read_bytes()
            report, issues = audit_file(path)
            self.assertEqual(report["recovered_complete_failures"], 1)
            self.assertEqual(issues[0]["candidate_decision"], {i:0.5 for i in (1,2,3,4)})
            self.assertEqual(path.read_bytes(), before)

    def test_audit_uses_exact_received_budget_instead_of_old_float_prompt(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "dialogue.json"
            turn = {"content": "{1:3.6,2:3.6,3:3.6,4:3.6}",
                    "prompt": "Neighbors sent: {1:1.2,2:1.2,3:1.2,4:1.2}, you received:14.399999999999999.",
                    "executed_decision": {str(i):3.6 for i in (1,2,3,4)},
                    "parse_failed": False, "finish_reason": "stop"}
            atomic_json(path, {"dialogs":{"round_1":{"agent_0_free":{"as_trustee":turn}}}})
            report, issues = audit_file(path)
            self.assertEqual(report["changed_complete_decisions"], 0)
            self.assertEqual(issues, [])
    def test_legacy_byte_fingerprint_and_changed_source_rejection(self):
        with tempfile.TemporaryDirectory(dir=ROOT / "outputs") as temp:
            out = Path(temp) / "fingerprint"
            self.assertEqual(main(["--mock", "--output-dir", str(out)]), 0)
            path = out / "checkpoint.json"
            saved = json.loads(path.read_text(encoding="utf-8"))
            saved["engine_sha256"] = hashlib.sha256(b"".join(
                (ROOT / "src/experiment" / name).read_bytes() for name in CORE_FILES)).hexdigest()
            atomic_json(path, saved)
            self.assertEqual(main(["--resume", str(out)]), 0)
            saved["engine_sha256"] = "0" * 64
            atomic_json(path, saved)
            with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as raised:
                main(["--resume", str(out)])
            self.assertEqual(raised.exception.code, 2)

    def test_batch_resume_and_seed_export_precision(self):
        with tempfile.TemporaryDirectory(dir=ROOT / "outputs") as temp:
            config = Path(temp) / "batch.json"
            config.write_text(json.dumps({"model_type_list": ["model-a", "model-b"]}), encoding="utf-8")
            out = Path(temp) / "batch"
            gateways = []
            def factory(p, name, **kwargs):
                gateway = CountingMock(p, pause_at=2 if name == "model-b" else None)
                gateways.append(gateway)
                return gateway
            with patch("src.experiment.run_main.MockChat", side_effect=factory):
                self.assertEqual(main(["--mock", "--config", str(config), "--output-dir", str(out)]), 2)
            saved = json.loads((out / "checkpoint.json").read_text(encoding="utf-8"))
            self.assertEqual((saved["scenario_index"], len(saved["completed"]), saved["current"]["next_agent"]), (1, 1, 2))
            self.assertEqual(main(["--resume", str(out)]), 0)
            import pandas as pd
            from openpyxl import load_workbook
            df = pd.read_csv(out / "agent_rounds.csv", dtype={"seed": "string"})
            self.assertEqual(len(df), 18)
            seeds = {str(s["seed"]) for s in saved["scenarios"]}
            self.assertEqual(set(df["seed"]), seeds)
            workbook = load_workbook(out / "pair_data.xlsx")
            sheet = workbook.active
            seed_col = next(cell.column for cell in sheet[1] if cell.value == "seed")
            values = {sheet.cell(row, seed_col).value for row in range(2, sheet.max_row + 1)}
            self.assertEqual(values, seeds)
            workbook.close()

    def test_mock_export_resume_and_context_pause(self):
        with tempfile.TemporaryDirectory(dir=ROOT / "outputs") as temp:
            out=Path(temp)/"run"
            self.assertEqual(main(["--mock","--output-dir",str(out)]),0)
            saved=json.loads((out/"checkpoint.json").read_text(encoding="utf-8"))
            self.assertEqual(saved["status"],"complete")
            self.assertEqual(saved["completed"][0]["summary"]["accepted_calls"],18)
            event_lines=len((out/"events.jsonl").read_text(encoding="utf-8").splitlines())
            self.assertEqual(main(["--resume",str(out)]),0)
            self.assertEqual(len((out/"events.jsonl").read_text(encoding="utf-8").splitlines()),event_lines)
            config=Path(temp)/"small_context.json"
            config.write_text(json.dumps({"context_limit":600}),encoding="utf-8")
            out2=Path(temp)/"paused"
            self.assertEqual(main(["--mock","--config",str(config),"--output-dir",str(out2)]),2)
            saved=json.loads((out2/"checkpoint.json").read_text(encoding="utf-8"))
            self.assertEqual(saved["current"]["next_agent"],0)
            config.write_text(json.dumps({"context_limit":131072}),encoding="utf-8")
            self.assertEqual(main(["--resume",str(out2),"--config",str(config)]),0)

    def test_experiment_imports_without_mesa_camel_or_analysis(self):
        script='''import sys, importlib.abc
class Deny(importlib.abc.MetaPathFinder):
 def find_spec(self, fullname, path, target=None):
  if fullname.split('.')[0] in {'mesa','camel','mcp','tiktoken','matplotlib','seaborn','sklearn'}:
   raise AssertionError('unwanted dependency: '+fullname)
sys.meta_path.insert(0,Deny())
from src.experiment.run_main import main
raise SystemExit(main(['--dry-run']))
'''
        result=subprocess.run([sys.executable,"-c",script],cwd=ROOT,capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)

    def test_redaction_and_bad_configs(self):
        self.assertEqual(scrub({1:"secret", "api_key":"secret"},"secret"),{1:"[REDACTED]"})
        for config in ({"num_agents":4},{"proportions":0.5},{"history_rounds":0},{"seed":True},
                       {"request_timeout":0},{"game_type":"pd_game"},{"max_retries":-1}):
            with self.assertRaises(ValueError):
                params().update(config)


if __name__ == "__main__":
    unittest.main()
